"""Spreadsheet operations built directly on openpyxl.

Every function loads the workbook from disk, performs its work, saves when it
mutates, and returns a JSON-serializable dict. Errors are raised as
``ExcelToolError`` and converted to structured results by ``tools.execute_tool``.
"""

from __future__ import annotations

from typing import Any, Dict, List

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import range_boundaries
from openpyxl.utils.cell import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from config import INSPECT_PREVIEW_ROWS


class ExcelToolError(Exception):
    """Raised for any expected, user-recoverable spreadsheet error."""


def _load(workbook_path: str, data_only: bool = False) -> openpyxl.Workbook:
    try:
        return openpyxl.load_workbook(workbook_path, data_only=data_only)
    except (InvalidFileException, OSError) as exc:
        raise ExcelToolError(f"Could not open workbook: {exc}") from exc


def _save(wb: openpyxl.Workbook, workbook_path: str) -> None:
    """Save, flagging a full recalculation on next open.

    openpyxl writes formula text but no cached result, so spreadsheet apps that
    trust the cache (LibreOffice, Numbers, Google Sheets) would otherwise show 0
    for freshly written formulas until a manual recalc.
    """
    wb.calculation.fullCalcOnLoad = True
    wb.save(workbook_path)


def _get_sheet(wb: openpyxl.Workbook, sheet: str):
    if sheet not in wb.sheetnames:
        raise ExcelToolError(
            f"Sheet '{sheet}' does not exist. Available sheets: {wb.sheetnames}"
        )
    return wb[sheet]


def _parse_range(cell_range: str):
    """Return (min_col, min_row, max_col, max_row) for an A1 range or single cell."""
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range.strip())
    except (ValueError, TypeError) as exc:
        raise ExcelToolError(
            f"Invalid cell range '{cell_range}'. Use A1 notation like D2 or D2:D100."
        ) from exc
    if None in (min_col, min_row, max_col, max_row):
        raise ExcelToolError(
            f"Invalid cell range '{cell_range}'. Whole-row/column ranges are not "
            f"supported; use bounded ranges like A1:C10."
        )
    return min_col, min_row, max_col, max_row


def _looks_like_header(row: List[Any]) -> bool:
    non_empty = [c for c in row if c is not None and str(c).strip() != ""]
    if not non_empty:
        return False
    return all(isinstance(c, str) for c in non_empty)


def inspect_workbook(workbook_path: str) -> Dict[str, Any]:
    wb = _load(workbook_path)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        preview: List[List[Any]] = []
        for row in ws.iter_rows(
            min_row=1, max_row=INSPECT_PREVIEW_ROWS, values_only=True
        ):
            preview.append(list(row))
        headers = None
        if preview and _looks_like_header(preview[0]):
            headers = [
                {"column": get_column_letter(i + 1), "name": v}
                for i, v in enumerate(preview[0])
                if v is not None and str(v).strip() != ""
            ]
        sheets.append(
            {
                "name": name,
                "dimensions": ws.dimensions,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "preview_rows": preview,
                "headers": headers,
            }
        )
    return {"success": True, "sheets": sheets}


def read_range(workbook_path: str, sheet: str, cell_range: str) -> Dict[str, Any]:
    wb = _load(workbook_path)
    ws = _get_sheet(wb, sheet)
    _parse_range(cell_range)

    wb_values = _load(workbook_path, data_only=True)
    ws_values = wb_values[sheet]

    cells: List[Dict[str, Any]] = []
    try:
        stored_rows = ws[cell_range]
        value_rows = ws_values[cell_range]
    except (ValueError, TypeError) as exc:
        raise ExcelToolError(f"Invalid cell range '{cell_range}': {exc}") from exc

    # Normalise single-cell access (openpyxl returns a bare Cell, not a tuple).
    if not isinstance(stored_rows, tuple):
        stored_rows = ((stored_rows,),)
        value_rows = ((value_rows,),)
    elif stored_rows and not isinstance(stored_rows[0], tuple):
        stored_rows = (stored_rows,)
        value_rows = (value_rows,)

    for stored_row, value_row in zip(stored_rows, value_rows):
        for stored_cell, value_cell in zip(stored_row, value_row):
            is_formula = (
                isinstance(stored_cell.value, str)
                and stored_cell.value.startswith("=")
            )
            cells.append(
                {
                    "coordinate": stored_cell.coordinate,
                    "value": value_cell.value if is_formula else stored_cell.value,
                    "formula": stored_cell.value if is_formula else None,
                }
            )
    return {"success": True, "sheet": sheet, "range": cell_range, "cells": cells}


def write_range(
    workbook_path: str, sheet: str, start_cell: str, values: List[List[Any]]
) -> Dict[str, Any]:
    if not isinstance(values, list) or not all(isinstance(r, list) for r in values):
        raise ExcelToolError("'values' must be a 2D array (list of lists).")

    wb = _load(workbook_path)
    ws = _get_sheet(wb, sheet)
    min_col, min_row, _, _ = _parse_range(start_cell)

    modified = 0
    for r, row in enumerate(values):
        for c, val in enumerate(row):
            ws.cell(row=min_row + r, column=min_col + c, value=val)
            modified += 1

    _save(wb, workbook_path)
    end_col = get_column_letter(min_col + (len(values[0]) - 1 if values else 0))
    end_ref = f"{end_col}{min_row + max(len(values) - 1, 0)}"
    return {
        "success": True,
        "message": f"Wrote {modified} cells to {sheet}!{start_cell}:{end_ref}",
        "cells_modified": modified,
    }


def write_formula(
    workbook_path: str, sheet: str, target_range: str, formula: str
) -> Dict[str, Any]:
    if not isinstance(formula, str) or not formula.startswith("="):
        raise ExcelToolError("'formula' must be a string starting with '='.")

    wb = _load(workbook_path)
    ws = _get_sheet(wb, sheet)
    min_col, min_row, max_col, max_row = _parse_range(target_range)
    origin = f"{get_column_letter(min_col)}{min_row}"

    modified = 0
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            dest = f"{get_column_letter(col)}{row}"
            if dest == origin:
                cell_formula = formula
            else:
                cell_formula = Translator(formula, origin=origin).translate_formula(
                    dest
                )
            ws[dest] = cell_formula
            modified += 1

    _save(wb, workbook_path)
    return {
        "success": True,
        "message": f"Wrote formula to {modified} cells in {sheet}!{target_range}",
        "cells_modified": modified,
    }


def create_sheet(workbook_path: str, sheet_name: str) -> Dict[str, Any]:
    wb = _load(workbook_path)
    if sheet_name in wb.sheetnames:
        raise ExcelToolError(f"Sheet '{sheet_name}' already exists.")
    wb.create_sheet(title=sheet_name)
    _save(wb, workbook_path)
    return {
        "success": True,
        "message": f"Created sheet '{sheet_name}'",
        "sheets": wb.sheetnames,
    }


def save_workbook(workbook_path: str) -> Dict[str, Any]:
    wb = _load(workbook_path)
    _save(wb, workbook_path)
    return {"success": True, "message": f"Saved workbook {workbook_path}"}
