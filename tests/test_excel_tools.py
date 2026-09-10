"""Direct tests of the openpyxl-backed tool functions."""

import openpyxl
import pytest

from excel_tools import (
    ExcelToolError,
    create_sheet,
    inspect_workbook,
    read_range,
    save_workbook,
    write_formula,
    write_range,
)


def test_inspect_workbook_reports_sheets_and_headers(workbook_path):
    result = inspect_workbook(workbook_path)
    names = [s["name"] for s in result["sheets"]]
    assert names == ["Sales", "Notes"]
    sales = result["sheets"][0]
    header_names = [h["name"] for h in sales["headers"]]
    assert header_names == ["Product", "Revenue", "Cost"]


def test_read_range_values_and_coordinates(workbook_path):
    result = read_range(workbook_path, "Sales", "A1:C1")
    assert [c["coordinate"] for c in result["cells"]] == ["A1", "B1", "C1"]
    assert [c["value"] for c in result["cells"]] == ["Product", "Revenue", "Cost"]


def test_read_single_cell(workbook_path):
    result = read_range(workbook_path, "Sales", "B2")
    assert result["cells"] == [
        {"coordinate": "B2", "value": 100, "formula": None}
    ]


def test_write_range_writes_and_saves(workbook_path):
    result = write_range(workbook_path, "Sales", "E1", [["Tag"], ["a"], ["b"]])
    assert result["cells_modified"] == 3
    wb = openpyxl.load_workbook(workbook_path)
    assert wb["Sales"]["E1"].value == "Tag"
    assert wb["Sales"]["E3"].value == "b"


def test_write_formula_fills_range_with_adjusted_refs(workbook_path):
    result = write_formula(workbook_path, "Sales", "D2:D4", "=(B2-C2)/B2")
    assert result["cells_modified"] == 3
    wb = openpyxl.load_workbook(workbook_path)
    assert wb["Sales"]["D2"].value == "=(B2-C2)/B2"
    assert wb["Sales"]["D3"].value == "=(B3-C3)/B3"
    assert wb["Sales"]["D4"].value == "=(B4-C4)/B4"


def test_write_formula_preserves_unrelated_data(workbook_path):
    write_formula(workbook_path, "Sales", "D2:D4", "=(B2-C2)/B2")
    wb = openpyxl.load_workbook(workbook_path)
    assert wb["Sales"]["A2"].value == "Widget"
    assert wb["Sales"]["B3"].value == 200


def test_write_formula_rejects_non_formula(workbook_path):
    with pytest.raises(ExcelToolError):
        write_formula(workbook_path, "Sales", "D2", "B2-C2")


def test_create_sheet(workbook_path):
    result = create_sheet(workbook_path, "Summary")
    assert "Summary" in result["sheets"]
    wb = openpyxl.load_workbook(workbook_path)
    assert "Summary" in wb.sheetnames


def test_create_sheet_rejects_duplicate(workbook_path):
    with pytest.raises(ExcelToolError):
        create_sheet(workbook_path, "Sales")


def test_missing_sheet_error(workbook_path):
    with pytest.raises(ExcelToolError) as exc:
        read_range(workbook_path, "Revenue", "A1:B2")
    assert "does not exist" in str(exc.value)


def test_invalid_cell_range_error(workbook_path):
    with pytest.raises(ExcelToolError):
        read_range(workbook_path, "Sales", "not-a-range")


def test_whole_column_range_rejected(workbook_path):
    with pytest.raises(ExcelToolError):
        write_formula(workbook_path, "Sales", "D:D", "=1")


def test_save_workbook(workbook_path):
    result = save_workbook(workbook_path)
    assert result["success"] is True
